#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Interroga l'API EUDAMED UDI per ogni REF presente in un file .xlsx
e produce un file Excel di risultato nella cartella del programma.

Dipendenze:  pip install requests openpyxl
"""

import sys
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

import requests
from openpyxl import Workbook, load_workbook

API_URL = (
    "https://api.datalake.sante.service.ec.europa.eu/eudamed/udi"
    "?REFERENCE={ref}&format=json&api-version=v1.0"
)
PAUSA_MINIMA = 0.5   # secondi tra una chiamata e la successiva
TIMEOUT = 30         # secondi di timeout per ogni chiamata

NON_TROVATO = "#NON TROVATO"
ERRORE = "#ERRORE"

INTESTAZIONI = [
    "PRODUTTORE", "EU_PRODUTTORE", "REF", "DESCRIZIONE",
    "EU_DESCRIZIONE", "EU_UDI", "EU_BASIC_UDI",
]


def cartella_programma() -> Path:
    if getattr(sys, "frozen", False):  # eseguibile creato con PyInstaller
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def scegli_file() -> Path:
    """Chiede all'utente di caricare il file .xlsx (finestra di dialogo o, in alternativa, da terminale)."""
    print("Seleziona il file .xlsx da elaborare "
          "(colonne: A=PRODUTTORE, B=REF, C=DESCRIZIONE, intestazione in prima riga)...")
    percorso = ""
    try:
        import tkinter as tk
        from tkinter import filedialog

        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        percorso = filedialog.askopenfilename(
            title="Carica il file Excel (.xlsx)",
            filetypes=[("File Excel", "*.xlsx")],
        )
        root.destroy()
    except Exception:
        percorso = input("Percorso del file .xlsx: ").strip().strip('"')

    if not percorso:
        print("Nessun file selezionato. Uscita.")
        sys.exit(1)
    return Path(percorso)


def come_testo(valore) -> str:
    """Converte il valore di una cella in testo (3451.0 -> '3451')."""
    if valore is None:
        return ""
    if isinstance(valore, float) and valore.is_integer():
        return str(int(valore))
    return str(valore).strip()


def leggi_righe(percorso: Path):
    wb = load_workbook(percorso, read_only=True, data_only=True)
    ws = wb.active
    righe = []
    for produttore, ref, descrizione, *_ in ws.iter_rows(
        min_row=2, max_col=3, values_only=True
    ):
        # salta le righe completamente vuote
        if produttore is None and ref is None and descrizione is None:
            continue
        righe.append((come_testo(produttore), come_testo(ref), come_testo(descrizione)))
    wb.close()
    return righe


class ClientApi:
    def __init__(self):
        self.sessione = requests.Session()
        self.ultima_chiamata = 0.0

    def cerca(self, ref: str):
        """Restituisce la lista 'value' dell'API. Solleva eccezione in caso di errore."""
        # rispetta la pausa minima tra una chiamata e l'altra
        trascorso = time.monotonic() - self.ultima_chiamata
        if trascorso < PAUSA_MINIMA:
            time.sleep(PAUSA_MINIMA - trascorso)

        try:
            risposta = self.sessione.get(API_URL.format(ref=quote(ref, safe="")), timeout=TIMEOUT)
        finally:
            self.ultima_chiamata = time.monotonic()

        risposta.raise_for_status()
        dati = risposta.json()
        valore = dati.get("value", [])
        return valore if isinstance(valore, list) else []


def salva_risultato(righe_out, cartella: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Risultato"
    ws.append(INTESTAZIONI)
    for r in righe_out:
        ws.append(r)

    larghezze = [28, 32, 20, 40, 40, 22, 28]
    for i, l in enumerate(larghezze):
        ws.column_dimensions[chr(ord("A") + i)].width = l
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # REF, UDI e BASIC UDI sempre come testo (evita perdita di zeri iniziali)
    for riga in ws.iter_rows(min_row=2):
        for idx in (2, 5, 6):
            riga[idx].number_format = "@"

    nome = f"risultato_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    destinazione = cartella / nome
    wb.save(destinazione)
    return destinazione


def main():
    file_input = scegli_file()
    try:
        righe = leggi_righe(file_input)
    except Exception as e:
        print(f"Impossibile leggere il file: {e}")
        sys.exit(1)

    totale = len(righe)
    if totale == 0:
        print("Il file non contiene righe da elaborare.")
        sys.exit(0)

    print(f"Righe da elaborare: {totale}\n")
    client = ClientApi()
    risultati = []

    try:
        for n, (produttore, ref, descrizione) in enumerate(righe, start=1):
            print(f"Chiamata {n} di {totale}  (REF: {ref or '-'})", end=" ... ", flush=True)

            if not ref:
                risultati.append([produttore, "", ref, descrizione, "", NON_TROVATO, ""])
                print("REF vuoto")
                continue

            try:
                trovati = client.cerca(ref)
            except Exception as e:
                risultati.append([produttore, "", ref, descrizione, "", ERRORE, ""])
                print(f"ERRORE ({e})")
                continue

            if not trovati:
                risultati.append([produttore, "", ref, descrizione, "", NON_TROVATO, ""])
                print("non trovato")
                continue

            for item in trovati:
                risultati.append([
                    produttore,
                    item.get("MF_NAME") or "",
                    ref,
                    descrizione,
                    item.get("TRADE_NAME") or "",
                    item.get("PRIMARY_DI") or "",
                    item.get("BASIC_UDI") or "",
                ])
            print(f"{len(trovati)} risultat{'o' if len(trovati) == 1 else 'i'}")

    except KeyboardInterrupt:
        print("\nInterrotto dall'utente: salvo i risultati parziali.")

    destinazione = salva_risultato(risultati, cartella_programma())
    print(f"\nCompletato. File di risultato salvato in:\n{destinazione}")


if __name__ == "__main__":
    main()
